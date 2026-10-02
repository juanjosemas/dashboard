# -*- coding: utf-8 -*-
"""Amplia 'Resumen Mensual por Empleado' a 18 meses (jul-2026 -> dic-2027)
para que los recibos de meses nuevos (oct-2026 en adelante) salgan solos."""
import openpyxl
from copy import copy

PATH = r"C:\Users\jjmax\Downloads\Control Gasolina\Control_Gasolina_Empresa.xlsx"
MESES = 18
FILA_INI = 4          # primera fila de mes
FILA_TOT_OLD = 7      # donde estaba TOTAL GENERAL
REG_A, REG_C, REG_G = "$A$5:$A$207", "$C$5:$C$207", "$G$5:$G$207"
FMT_OCULTA_CERO = '#,##0.00" €";-#,##0.00" €";""'

wb = openpyxl.load_workbook(PATH)
ws = wb["Resumen Mensual por Empleado"]

# guardar estilos originales antes de reescribir
st_mes_a = copy(ws.cell(row=FILA_INI, column=1)._style)
st_mes_b = copy(ws.cell(row=FILA_INI, column=2)._style)
st_tot_a = copy(ws.cell(row=FILA_TOT_OLD, column=1)._style)
st_tot_b = copy(ws.cell(row=FILA_TOT_OLD, column=2)._style)

FILA_TOT = FILA_INI + MESES          # nueva fila TOTAL GENERAL

for i in range(MESES):
    r = FILA_INI + i
    # mes
    a = ws.cell(row=r, column=1)
    a.value = "=DATE(2026,7,1)" if i == 0 else "=EDATE($A{},1)".format(r - 1)
    a._style = copy(st_mes_a)
    # importes por trabajador (cabeceras en la fila 3) y total del mes
    for col in range(2, 11):
        cl = openpyxl.utils.get_column_letter(col)
        c = ws.cell(row=r, column=col)
        c.value = ('=SUMIFS(\'Registro Gasolina\'!{g},'
                   '\'Registro Gasolina\'!{a},">="&$A{r},'
                   '\'Registro Gasolina\'!{a},"<"&EDATE($A{r},1),'
                   '\'Registro Gasolina\'!{c},{cl}$3)').format(
            g=REG_G, a=REG_A, c=REG_C, r=r, cl=cl)
        c._style = copy(st_mes_b)
        c.number_format = FMT_OCULTA_CERO
    k = ws.cell(row=r, column=11)
    k.value = "=SUM(B{r}:J{r})".format(r=r)
    k._style = copy(st_mes_b)
    k.number_format = FMT_OCULTA_CERO

# (la antigua fila TOTAL 7 cae dentro del rango de meses, se pisa sola arriba)

t = ws.cell(row=FILA_TOT, column=1, value="TOTAL GENERAL")
t._style = copy(st_tot_a)
for col in range(2, 12):
    cl = openpyxl.utils.get_column_letter(col)
    c = ws.cell(row=FILA_TOT, column=col,
                value="=SUM({cl}{a}:{cl}{b})".format(cl=cl, a=FILA_INI, b=FILA_TOT - 1))
    c._style = copy(st_tot_b)
    c.number_format = FMT_OCULTA_CERO

wb.calculation.fullCalcOnLoad = True
wb.save(PATH)
print("OK: filas de mes", FILA_INI, "->", FILA_TOT - 1, "| TOTAL en fila", FILA_TOT)
